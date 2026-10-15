import io
from typing import List, Any
from app.models.schemas import TailoredResume

def extract_bullet_text(bullet: Any) -> str:
    """Extract clean string from bullet (which may be a string, dict, or object)."""
    if not bullet:
        return ""
    if isinstance(bullet, dict):
        text = bullet.get("tailored") or bullet.get("original") or bullet.get("text") or bullet.get("desc") or bullet.get("description")
        if text:
            return str(text).strip()
        parts = [f"{k}: {v}" for k, v in bullet.items() if v and isinstance(v, (str, int, float))]
        return " | ".join(parts)
    return str(bullet).strip()

def format_education_entry(edu: Any, separator: str = " - ") -> str:
    """Format education item whether it's a dict or string."""
    if not edu:
        return ""
    if isinstance(edu, dict):
        deg = str(edu.get("degree") or edu.get("title") or "").strip()
        inst = str(edu.get("institution") or edu.get("college") or edu.get("school") or "").strip()
        dur = str(edu.get("duration") or edu.get("year") or edu.get("dates") or "").strip()
        parts = []
        if deg:
            parts.append(deg)
        if inst:
            parts.append(inst)
        if dur:
            parts.append(f"({dur})")
        return separator.join(parts) if parts else "Education qualification"
    return str(edu).strip()

def build_contact_line(resume: TailoredResume, separator: str = " | ") -> str:
    """Build a clean single line of contact information."""
    parts = []
    if resume.email:
        parts.append(str(resume.email).strip())
    if resume.location:
        parts.append(str(resume.location).strip())
    if resume.linkedin:
        parts.append(str(resume.linkedin).strip())
    if not parts and resume.header_tags:
        for t in resume.header_tags:
            if t and isinstance(t, str) and t.strip():
                parts.append(t.strip())
    return separator.join(parts)

def safe_reportlab_text(text: Any) -> str:
    """Escapes XML entities so ReportLab Paragraph never encounters invalid XML."""
    if text is None:
        return ""
    s = str(text)
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def parse_entry_helper(entry: Any):
    """Normalize experience entry whether it's a dict, object, or string."""
    if isinstance(entry, dict):
        title = str(entry.get("title") or "").strip()
        company = str(entry.get("company") or "").strip()
        dur = str(entry.get("duration") or "").strip()
        bullets = entry.get("bullets", [])
    elif hasattr(entry, "title"):
        title = str(getattr(entry, "title", "") or "").strip()
        company = str(getattr(entry, "company", "") or "").strip()
        dur = str(getattr(entry, "duration", "") or "").strip()
        bullets = getattr(entry, "bullets", [])
    else:
        title = str(entry).strip()
        company = ""
        dur = ""
        bullets = []
    return title, company, dur, bullets

def generate_txt_resume(resume: TailoredResume) -> bytes:
    """Generate a clean, structured plain text resume."""
    output = []
    cand_name = (resume.candidate_name or "Candidate").upper()
    target = resume.target_role or "Professional"
    
    output.append("=" * 64)
    output.append(f"  {cand_name}")
    output.append(f"  {target}")
    contact = build_contact_line(resume)
    if contact:
        output.append(f"  {contact}")
    output.append("=" * 64)
    output.append("")
    
    # Career Objective
    if resume.career_objective:
        output.append("CAREER OBJECTIVE")
        output.append("-" * 32)
        output.append(str(resume.career_objective).strip())
        output.append("")
        
    # Professional Summary
    if resume.professional_summary:
        output.append("PROFESSIONAL SUMMARY")
        output.append("-" * 32)
        output.append(str(resume.professional_summary).strip())
        output.append("")
        
    # Matched Skills
    skills = resume.matched_skills or resume.all_skills
    if skills:
        output.append("CORE SKILLS & COMPETENCIES")
        output.append("-" * 32)
        output.append(", ".join([str(s).strip() for s in skills if s]))
        output.append("")
        
    # Experience Entries
    if resume.experience_entries:
        output.append("PROFESSIONAL EXPERIENCE")
        output.append("-" * 32)
        for entry in resume.experience_entries:
            title, company, dur, bullets = parse_entry_helper(entry)
            header_parts = [p for p in [title, company, dur] if p]
            if header_parts:
                output.append(" | ".join(header_parts))
            for b in bullets:
                clean_b = extract_bullet_text(b)
                if clean_b:
                    output.append(f"  - {clean_b}")
            output.append("")
    elif resume.experience_highlights:
        output.append("PROFESSIONAL EXPERIENCE")
        output.append("-" * 32)
        for bullet in resume.experience_highlights:
            clean_b = extract_bullet_text(bullet)
            if clean_b:
                output.append(f"  - {clean_b}")
        output.append("")
        
    # Education
    if resume.education:
        output.append("EDUCATION")
        output.append("-" * 32)
        for edu in resume.education:
            clean_edu = format_education_entry(edu)
            if clean_edu:
                output.append(f"  - {clean_edu}")
        output.append("")
        
    # Additional Information / Projects
    if resume.additional_information:
        output.append("PROJECTS & TECHNICAL DOMAINS")
        output.append("-" * 32)
        for item in resume.additional_information:
            clean_item = extract_bullet_text(item)
            if clean_item:
                output.append(f"  - {clean_item}")
        output.append("")
        
    text_content = "\n".join(output)
    return text_content.encode("utf-8")

def generate_pdf_resume(resume: TailoredResume) -> bytes:
    """Generate an elegant, professional PDF document using reportlab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=17,
            leading=21,
            textColor=colors.HexColor('#064e3b')
        )
        
        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10.5,
            leading=14,
            textColor=colors.HexColor('#047857')
        )
        
        meta_style = ParagraphStyle(
            'MetaTags',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor('#4b5563')
        )
        
        heading_style = ParagraphStyle(
            'SectionHeading',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#111827'),
            spaceBefore=8,
            spaceAfter=3
        )
        
        entry_title_style = ParagraphStyle(
            'EntryTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#1f2937'),
            spaceBefore=4,
            spaceAfter=1
        )
        
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#374151')
        )
        
        bullet_style = ParagraphStyle(
            'Bullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=12.5,
            textColor=colors.HexColor('#374151'),
            leftIndent=10
        )
        
        story = []
        
        # Candidate Header
        story.append(Paragraph(safe_reportlab_text(resume.candidate_name or "Candidate"), title_style))
        story.append(Spacer(1, 2))
        story.append(Paragraph(safe_reportlab_text(resume.target_role or "Professional"), subtitle_style))
        story.append(Spacer(1, 2))
        
        contact = build_contact_line(resume)
        if contact:
            story.append(Paragraph(safe_reportlab_text(contact), meta_style))
            story.append(Spacer(1, 4))
            
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#d1d5db'), spaceAfter=6))
        
        # Professional Summary
        if resume.professional_summary:
            story.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
            story.append(Paragraph(safe_reportlab_text(resume.professional_summary), body_style))
            story.append(Spacer(1, 5))
            
        # Matched Skills
        skills = resume.matched_skills or resume.all_skills
        if skills:
            story.append(Paragraph("TECHNICAL COMPETENCIES & MATCHED SKILLS", heading_style))
            clean_skills = [safe_reportlab_text(str(s).strip()) for s in skills if s]
            story.append(Paragraph(" &bull; ".join(clean_skills), body_style))
            story.append(Spacer(1, 5))
            
        # Experience
        if resume.experience_entries:
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", heading_style))
            for entry in resume.experience_entries:
                title, company, dur, bullets = parse_entry_helper(entry)
                header_text = f"<b>{safe_reportlab_text(title)}</b>"
                if company:
                    header_text += f" &mdash; <i>{safe_reportlab_text(company)}</i>"
                if dur:
                    header_text += f" <font color='#6b7280'>({safe_reportlab_text(dur)})</font>"
                story.append(Paragraph(header_text, entry_title_style))
                
                for b in bullets:
                    clean_b = extract_bullet_text(b)
                    if clean_b:
                        story.append(Paragraph(f"&bull; {safe_reportlab_text(clean_b)}", bullet_style))
                story.append(Spacer(1, 4))
            story.append(Spacer(1, 3))
        elif resume.experience_highlights:
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", heading_style))
            for item in resume.experience_highlights[:12]:
                clean_b = extract_bullet_text(item)
                if clean_b:
                    story.append(Paragraph(f"&bull; {safe_reportlab_text(clean_b)}", bullet_style))
            story.append(Spacer(1, 5))
            
        # Education
        if resume.education:
            story.append(Paragraph("EDUCATION", heading_style))
            for edu in resume.education:
                formatted_edu = format_education_entry(edu)
                if formatted_edu:
                    story.append(Paragraph(f"&bull; {safe_reportlab_text(formatted_edu)}", bullet_style))
            story.append(Spacer(1, 5))
            
        # Additional Info / Projects
        if resume.additional_information:
            story.append(Paragraph("PROJECTS & ADDITIONAL INFORMATION", heading_style))
            for info in resume.additional_information[:8]:
                clean_info = extract_bullet_text(info)
                if clean_info:
                    story.append(Paragraph(f"&bull; {safe_reportlab_text(clean_info)}", bullet_style))
                    
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
    except Exception as e:
        print(f"ReportLab PDF generation error: {e}")
        return generate_txt_resume(resume)

def generate_docx_resume(resume: TailoredResume) -> bytes:
    """Generate a formatted Microsoft Word DOCX resume."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        
        doc = Document()
        
        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)
            
        # Title
        title_p = doc.add_paragraph()
        run = title_p.add_run(resume.candidate_name or "Candidate")
        run.bold = True
        run.font.size = Pt(17)
        run.font.color.rgb = RGBColor(6, 78, 59)
        title_p.paragraph_format.space_after = Pt(2)
        
        # Subtitle
        sub_p = doc.add_paragraph()
        s_run = sub_p.add_run(resume.target_role or "Professional")
        s_run.bold = True
        s_run.font.size = Pt(11)
        s_run.font.color.rgb = RGBColor(4, 120, 87)
        sub_p.paragraph_format.space_after = Pt(2)
        
        # Contact
        contact = build_contact_line(resume)
        if contact:
            meta_p = doc.add_paragraph()
            m_run = meta_p.add_run(contact)
            m_run.font.size = Pt(9)
            m_run.font.color.rgb = RGBColor(100, 116, 139)
            meta_p.paragraph_format.space_after = Pt(8)
            
        # Professional Summary
        if resume.professional_summary:
            doc.add_heading('PROFESSIONAL SUMMARY', level=2)
            doc.add_paragraph(resume.professional_summary)
            
        # Matched Skills
        skills = resume.matched_skills or resume.all_skills
        if skills:
            doc.add_heading('TECHNICAL COMPETENCIES & MATCHED SKILLS', level=2)
            clean_skills = [str(s).strip() for s in skills if s]
            doc.add_paragraph(" • ".join(clean_skills))
            
        # Experience
        if resume.experience_entries:
            doc.add_heading('PROFESSIONAL EXPERIENCE', level=2)
            for entry in resume.experience_entries:
                title, company, dur, bullets = parse_entry_helper(entry)
                p_entry = doc.add_paragraph()
                title_run = p_entry.add_run(title)
                title_run.bold = True
                if company:
                    comp_run = p_entry.add_run(f" — {company}")
                    comp_run.italic = True
                if dur:
                    dur_run = p_entry.add_run(f" ({dur})")
                    dur_run.font.color.rgb = RGBColor(100, 116, 139)
                p_entry.paragraph_format.space_after = Pt(2)
                
                for b in bullets:
                    clean_b = extract_bullet_text(b)
                    if clean_b:
                        doc.add_paragraph(clean_b, style='List Bullet')
        elif resume.experience_highlights:
            doc.add_heading('PROFESSIONAL EXPERIENCE', level=2)
            for bullet in resume.experience_highlights:
                clean_b = extract_bullet_text(bullet)
                if clean_b:
                    doc.add_paragraph(clean_b, style='List Bullet')
                
        # Education
        if resume.education:
            doc.add_heading('EDUCATION', level=2)
            for edu in resume.education:
                formatted_edu = format_education_entry(edu)
                if formatted_edu:
                    doc.add_paragraph(formatted_edu, style='List Bullet')
                
        # Additional Information
        if resume.additional_information:
            doc.add_heading('PROJECTS & TECHNICAL DOMAINS', level=2)
            for item in resume.additional_information:
                clean_item = extract_bullet_text(item)
                if clean_item:
                    doc.add_paragraph(clean_item, style='List Bullet')
                
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    except Exception as e:
        print(f"python-docx error: {e}")
        return generate_txt_resume(resume)
